"""Выгрузка истории дневных отчётов (data/daily_stats_history.json) в .xlsx (пункт D ТЗ)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.i18n import t
from app.daily_report import render_recommendation as render_daily_recommendation
from app.past_campaigns_report import render_recommendation as render_past_recommendation

HEADER_FILL = PatternFill(start_color="1A1A1E", end_color="1A1A1E", fill_type="solid")
HEADER_FONT = Font(color="BD94EB", bold=True)


def _columns():
    # Функция, а не модульная константа — t() читает язык из request.headers, доступного
    # только внутри запроса (все build_*_excel_bytes() вызываются live из routes/reports.py).
    return [
        (t("excel.col.date"), 12),
        (t("excel.col.campaign"), 32),
        (t("excel.col.objective"), 22),
        (t("excel.col.impressions"), 12),
        (t("excel.col.reach"), 12),
        (t("excel.col.ctr"), 10),
        (t("excel.col.cpm"), 10),
        (t("excel.col.spend"), 12),
        (t("excel.col.result"), 20),
        (t("excel.col.result_count"), 16),
        (t("excel.col.cost_per_result"), 16),
        (t("excel.col.roas"), 10),
        (t("excel.col.recommendation"), 50),
    ]


def _fmt(value):
    return round(value, 2) if isinstance(value, (int, float)) else None


def build_excel_bytes(history: list) -> bytes:
    """history — список снимков от app.daily_report.load_history() (новые первыми)."""
    wb = Workbook()
    ws = wb.active
    ws.title = t("excel.sheet.daily_stats")

    for col_idx, (title, width) in enumerate(_columns(), start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    row_idx = 2
    # history приходит от новых к старым — в таблице удобнее хронологический порядок
    for day in sorted(history, key=lambda d: d.get("date", "")):
        campaigns = day.get("campaigns") or []
        if not campaigns:
            note = render_daily_recommendation(day.get("note_code")) if day.get("note_code") else day.get("note")
            ws.cell(row=row_idx, column=1, value=day.get("date"))
            ws.cell(row=row_idx, column=2, value=note or t("reports.msg.no_active_campaigns"))
            row_idx += 1
            continue

        for c in campaigns:
            metrics = c.get("metrics") or {}
            result = c.get("result") or {}
            recommendation = (
                render_daily_recommendation(c.get("recommendation_code"), c.get("recommendation_params"))
                if c.get("recommendation_code") else c.get("recommendation")
            )
            ws.cell(row=row_idx, column=1, value=day.get("date"))
            ws.cell(row=row_idx, column=2, value=c.get("name"))
            ws.cell(row=row_idx, column=3, value=c.get("objective_label"))
            ws.cell(row=row_idx, column=4, value=_fmt(metrics.get("impressions")))
            ws.cell(row=row_idx, column=5, value=_fmt(metrics.get("reach")))
            ws.cell(row=row_idx, column=6, value=_fmt(metrics.get("ctr")))
            ws.cell(row=row_idx, column=7, value=_fmt(metrics.get("cpm")))
            ws.cell(row=row_idx, column=8, value=_fmt(metrics.get("spend")))
            ws.cell(row=row_idx, column=9, value=result.get("label") or t("common.no_data"))
            ws.cell(row=row_idx, column=10, value=_fmt(result.get("value")))
            ws.cell(row=row_idx, column=11, value=_fmt(result.get("cost_per_result")))
            ws.cell(row=row_idx, column=12, value=_fmt(result.get("roas")))
            ws.cell(row=row_idx, column=13, value=recommendation)
            row_idx += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _past_campaigns_columns():
    return [
        (t("excel.col.generated_at"), 17),
        (t("excel.col.period"), 12),
        (t("excel.col.filter"), 10),
        (t("excel.col.campaign"), 32),
        (t("excel.col.status"), 14),
        (t("excel.col.objective"), 22),
        (t("excel.col.impressions"), 12),
        (t("excel.col.reach"), 12),
        (t("excel.col.ctr"), 10),
        (t("excel.col.cpm"), 10),
        (t("excel.col.spend"), 12),
        (t("excel.col.result"), 20),
        (t("excel.col.result_count"), 16),
        (t("excel.col.cost_per_result"), 16),
        (t("excel.col.roas"), 10),
        (t("excel.col.recommendation"), 50),
    ]


def _period_label(period: str) -> str:
    key = f"ads.period.{period}"
    label = t(key)
    if label != key:
        return label
    return t(f"reports.past.period.{period}")


def _status_filter_label(status_filter: str) -> str:
    return t(f"reports.past.status.{status_filter}")


def build_past_campaigns_excel_bytes(history: list) -> bytes:
    """history — список снимков от app.past_campaigns_report.load_past_campaigns_history()
    (новые первыми). В отличие от дневного отчёта, каждый снимок здесь — отдельный запуск
    отчёта (может быть несколько за один день с разными периодами/фильтрами), поэтому
    сортировка по generated_at, а не по календарной дате."""
    wb = Workbook()
    ws = wb.active
    ws.title = t("excel.sheet.past_campaigns")

    for col_idx, (title, width) in enumerate(_past_campaigns_columns(), start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    row_idx = 2
    for snapshot in sorted(history, key=lambda s: s.get("generated_at", "")):
        generated_at = snapshot.get("generated_at", "")
        generated_label = generated_at[:16].replace("T", " ") if generated_at else ""
        period_label = _period_label(snapshot.get("period", ""))
        filter_label = _status_filter_label(snapshot.get("status_filter"))
        campaigns = snapshot.get("campaigns") or []

        if not campaigns:
            ws.cell(row=row_idx, column=1, value=generated_label)
            ws.cell(row=row_idx, column=2, value=period_label)
            ws.cell(row=row_idx, column=3, value=filter_label)
            ws.cell(row=row_idx, column=4, value=t("reports.past.no_campaigns_for_filter"))
            row_idx += 1
            continue

        for c in campaigns:
            metrics = c.get("metrics") or {}
            result = c.get("result") or {}
            recommendation = (
                render_past_recommendation(c.get("recommendation_code"), c.get("recommendation_params"))
                if c.get("recommendation_code") else c.get("recommendation")
            )
            ws.cell(row=row_idx, column=1, value=generated_label)
            ws.cell(row=row_idx, column=2, value=period_label)
            ws.cell(row=row_idx, column=3, value=filter_label)
            ws.cell(row=row_idx, column=4, value=c.get("name"))
            ws.cell(row=row_idx, column=5, value=c.get("effective_status") or c.get("status"))
            ws.cell(row=row_idx, column=6, value=c.get("objective_label"))
            ws.cell(row=row_idx, column=7, value=_fmt(metrics.get("impressions")))
            ws.cell(row=row_idx, column=8, value=_fmt(metrics.get("reach")))
            ws.cell(row=row_idx, column=9, value=_fmt(metrics.get("ctr")))
            ws.cell(row=row_idx, column=10, value=_fmt(metrics.get("cpm")))
            ws.cell(row=row_idx, column=11, value=_fmt(metrics.get("spend")))
            ws.cell(row=row_idx, column=12, value=result.get("label") or t("common.no_data"))
            ws.cell(row=row_idx, column=13, value=_fmt(result.get("value")))
            ws.cell(row=row_idx, column=14, value=_fmt(result.get("cost_per_result")))
            ws.cell(row=row_idx, column=15, value=_fmt(result.get("roas")))
            ws.cell(row=row_idx, column=16, value=recommendation)
            row_idx += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
